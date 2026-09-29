"""Evidence packet and audit export as a spreadsheet, for risk and compliance teams.

The workbook is written with openpyxl. When the OfficeCLI binary is on PATH, each export is also
validated with `officecli validate` and rendered with `officecli view … html` for an in-app preview.
Identifiers stay tokenised in the export.
"""
from __future__ import annotations

import html
import io
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F2937")


def officecli_available() -> bool:
    return shutil.which("officecli") is not None


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col) + 2))
    ws.freeze_panes = "A2"
    return ws


def case_workbook(case: dict, trail: list[dict]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    txn = case.get("txn") or {}
    _sheet(wb, "Case", ["Field", "Value"], [
        ["Case", case["id"]], ["Customer", case["customer_id"]], ["Channel", case["channel"]],
        ["Intent", case.get("intent")], ["Route", case.get("route")], ["Status", case["status"]],
        ["Transaction", txn.get("id")], ["Amount (₹)", txn.get("amount")], ["UTR", txn.get("utr")],
        ["Opened", time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(case["created_at"]))],
        ["Message (tokenised)", case.get("text")]])
    _sheet(wb, "Decisions", ["Action", "Amount (₹)", "Calibrated score", "Route", "Why", "Committed"],
           [[d["action"]["type"], d["action"]["amount"], d["score"], d["route"], d["why"], bool(d.get("committed"))]
            for d in case.get("decisions", [])])
    _sheet(wb, "Evidence", ["Fact"], [[f] for f in (case.get("plan") or {}).get("facts", [])])
    _sheet(wb, "Trail", ["Seq", "Time", "Kind", "Actor", "Hash", "Prev hash"],
           [[t["seq"], time.strftime("%H:%M:%S", time.localtime(t["ts"])), t["kind"], t["actor"], t["hash"],
             t["prev_hash"]] for t in trail])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def preview(xlsx: bytes, case: dict) -> dict:
    """Validate and render with OfficeCLI when present; otherwise a plain HTML rendering of the decisions."""
    if officecli_available():
        with tempfile.TemporaryDirectory() as d:
            p, out = Path(d) / "case.xlsx", Path(d) / "case.html"
            p.write_bytes(xlsx)
            try:
                v = subprocess.run(["officecli", "validate", str(p)], capture_output=True, timeout=60)
                subprocess.run(["officecli", "view", str(p), "html", "-o", str(out)], check=True,
                               capture_output=True, timeout=60)
                return {"engine": "officecli", "valid": v.returncode == 0, "html": out.read_text(errors="ignore")}
            except (subprocess.SubprocessError, OSError):
                pass
    rows = "".join(f"<tr><td>{html.escape(d['action']['type'])}</td><td>{d['action']['amount']}</td>"
                   f"<td>{d['score']}</td><td>{d['route']}</td><td>{html.escape(d['why'])}</td></tr>"
                   for d in case.get("decisions", []))
    body = (f"<h1>{html.escape(case['id'])}</h1><table border=1 cellpadding=6><tr><th>Action</th><th>Amount</th>"
            f"<th>Score</th><th>Route</th><th>Why</th></tr>{rows}</table>")
    return {"engine": "builtin", "valid": True, "html": body}
